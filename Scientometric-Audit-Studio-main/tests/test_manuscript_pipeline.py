"""
Automated Unit Tests for PDF/DOCX Extraction, Citation Style Detection,
and In-Text Citation Consistency Auditing.
"""
import io
import pytest
from docx import Document

from src.parsing.document_extractor import (
    extract_text_from_docx,
    extract_document,
    clean_line_hyphenation,
)
from src.parsing.section_segmenter import segment_manuscript
from src.parsing.citation_style_detector import (
    CitationStyleDetector,
    CitationConsistencyAuditor,
    CitationStyleFamily,
    unroll_numeric_range,
)
from src.parsing.reference_splitter import ReferenceSplitter


def test_unroll_numeric_range():
    assert unroll_numeric_range("1-4, 7, 9-11") == [1, 2, 3, 4, 7, 9, 10, 11]
    assert unroll_numeric_range("5–8") == [5, 6, 7, 8]
    assert unroll_numeric_range("12") == [12]
    assert unroll_numeric_range("") == []


def test_clean_line_hyphenation():
    raw = "This is a bio-\ntechnology study on sciento-\nmetrics."
    cleaned = clean_line_hyphenation(raw)
    assert "biotechnology" in cleaned
    assert "scientometrics" in cleaned


def test_citation_style_detector_numeric():
    body = (
        "Recent advances in deep learning [1] have transformed natural language processing [2, 3]. "
        "Furthermore, prior studies [4-6] demonstrated significant empirical gains. "
        "However, as noted in [8], computational constraints remain."
    )
    style, citations = CitationStyleDetector.scan_in_text_citations(body)
    assert style == CitationStyleFamily.NUMERIC_BRACKETED
    assert len(citations) == 4

    # Check unrolling: [2, 3] -> [2, 3], [4-6] -> [4, 5, 6]
    all_keys = set()
    for c in citations:
        all_keys.update(c["unrolled_keys"])
    assert all_keys == {1, 2, 3, 4, 5, 6, 8}


def test_citation_style_detector_author_year():
    body = (
        "Scholarly reference auditing has gained traction (Smith, 2020). "
        "According to Johnson et al. (2021), automated citation verification reduces errors. "
        "Earlier findings also support this hypothesis (Davis & Taylor, 2019; Williams, 2018)."
    )
    style, citations = CitationStyleDetector.scan_in_text_citations(body)
    assert style == CitationStyleFamily.AUTHOR_YEAR
    assert len(citations) >= 3


def test_consistency_auditor_numeric():
    # Cited in text: 1, 2, 3, 4, 5, 6, 8 (missing 7!)
    body = "Text citing [1], [2-3], [4-6], and [8]."
    style, citations = CitationStyleDetector.scan_in_text_citations(body)

    # Bibliography has 1..6, 8, 9 (orphan 9!)
    bibliography_items = [{"reference_no": i} for i in [1, 2, 3, 4, 5, 6, 8, 9]]

    report = CitationConsistencyAuditor.audit_consistency(style, citations, bibliography_items)
    assert report["orphan_count"] == 1
    assert 9 in report["orphan_reference_numbers"]
    assert report["missing_count"] == 0
    assert 7 in report["sequence_gaps"]
    assert report["consistency_score"] < 100.0


def test_section_segmenter_with_appendix():
    text = (
        "Introduction\n"
        "This is the introduction to our research paper on automated reference validation.\n\n"
        "Methodology\n"
        "We evaluate citations across several international publishing repositories.\n\n"
        "Results\n"
        "The experiments show high accuracy in resolving DOIs and validating titles.\n\n"
        "References\n"
        "[1] Smith J., Deep Learning in Scientometrics, Nature, 2020.\n"
        "[2] Brown A., Citation Analysis, Science, 2019.\n\n"
        "Appendix A: Supplementary Derivations\n"
        "Here are the mathematical proofs and additional figures."
    )
    body, refs, meta = segment_manuscript(text)
    assert meta["found_references_heading"] is True
    assert "[1] Smith J." in refs
    assert "[2] Brown A." in refs
    # Appendix must be stripped from references section!
    assert "Supplementary Derivations" not in refs
    assert meta["trailing_section_stripped"] is True


def test_section_segmenter_strips_spaced_small_caps_appendix_heading():
    text = (
        "Introduction\n" + ("Body text. " * 80) + "\n"
        "REFERENCES\n"
        "Ada Researcher and Grace Scientist. A reliable paper. Journal, 2020.\n"
        "A L ARGE LANGUAGE MODELS STILL NEED PARAMETER UPDATES\n"
        "Appendix discussion must not become a reference."
    )

    _body, references, metadata = segment_manuscript(text)

    assert "A reliable paper" in references
    assert "Appendix discussion" not in references
    assert metadata["trailing_section_stripped"] is True


def test_section_segmenter_strips_short_spaced_small_caps_appendix_heading():
    text = (
        "Introduction\n" + ("Body text. " * 80) + "\n"
        "REFERENCES\n"
        "Ligeng Zhu, Zhijian Liu, and Song Han. Deep leakage from gradients. "
        "NeurIPS, 2019.\n"
        "Yanlin Zhou, George Pu, Xiyao Ma, Xiaolin Li, and Dapeng Wu. "
        "Distilled one-shot federated learning. arXiv:2009.07999, 2020.\n"
        "A I MPLEMENTATION DETAILS\n"
        "Dataset condensation experiments involve six hyperparameters."
    )

    _body, references, metadata = segment_manuscript(text)

    assert "Deep leakage from gradients" in references
    assert "Dataset condensation experiments" not in references
    assert metadata["trailing_section_stripped"] is True


def test_docx_document_extraction():
    # Create an in-memory test DOCX
    doc = Document()
    doc.add_heading("Automated Scientometric Audit Framework", level=1)
    doc.add_paragraph("This paper investigates bibliographic reference errors in scholarly journals [1].")
    doc.add_paragraph("Following prior benchmarks [2, 3], we assess metadata alignment.")
    doc.add_heading("References", level=1)
    doc.add_paragraph("[1] LeCun Y., Bengio Y., Hinton G. Deep learning. Nature, 521, 436-444, (2015).")
    doc.add_paragraph("[2] Vaswani A., et al. Attention is all you need. NeurIPS, (2017).")
    doc.add_paragraph("[3] Devlin J., et al. BERT: Pre-training of deep bidirectional transformers. (2018).")

    buf = io.BytesIO()
    doc.save(buf)
    file_bytes = buf.getvalue()

    doc_content = extract_text_from_docx(file_bytes, filename="test_sample.docx")
    assert doc_content.file_type == "docx"
    assert "Automated Scientometric Audit Framework" in doc_content.full_text
    assert doc_content.word_count > 20

    body, refs, meta = segment_manuscript(doc_content.full_text)
    assert meta["found_references_heading"] is True
    assert "[1] LeCun Y." in refs

    ref_items = ReferenceSplitter.split_manuscript_references(refs)
    assert len(ref_items) == 3

    style, citations = CitationStyleDetector.scan_in_text_citations(body)
    assert style == CitationStyleFamily.NUMERIC_BRACKETED
    assert len(citations) == 2


def test_end_to_end_docx_custom_audit():
    # Create an in-memory manuscript
    doc = Document()
    doc.add_heading("Empirical Verification of Scientific Bibliographies", level=1)
    doc.add_paragraph("Academic bibliographies frequently suffer from swapped and phantom DOIs [1].")
    doc.add_paragraph("Modern transformer architectures [2] provide powerful text embedding representations.")
    doc.add_heading("References", level=1)
    doc.add_paragraph("[1] Aghion P Howitt P A model of growth through creative destruction Econometrica 1992 60 2 323 351 10.2307/2951599")
    doc.add_paragraph("[2] Vaswani A Attention is all you need Advances in Neural Information Processing Systems 2017")

    buf = io.BytesIO()
    doc.save(buf)
    file_bytes = buf.getvalue()

    from src.pipeline.custom_runner import run_custom_audit
    job = run_custom_audit(uploaded_file_bytes=file_bytes, uploaded_filename="sample_paper.docx")

    assert job["document_diagnostics"] is not None
    diag = job["document_diagnostics"]
    assert diag["citation_style"] == "NUMERIC_BRACKETED"
    assert diag["total_in_text_citations"] == 2
    assert diag["total_bibliography_references"] == 2
    assert diag["orphan_count"] == 0
    assert diag["missing_count"] == 0
    assert diag["consistency_score"] == 100.0

    assert len(job["items"]) == 2
    ref1 = job["items"][0]
    assert ref1["normalized_doi"] == "10.2307/2951599"
    assert ref1["doi_resolves"] is True


def test_split_inline_bracketed_references_and_control_chars():
    raw_blob = (
        "[1] Kumar, R., Singh, J., Singh, B., & Rana, M. K. \"Usability of OPAC in Libraries.\"\x0c "
        "Library Philosophy and Practice, 2018. doi:10.1177/0165551520952315. "
        "Some trailing text\x00[22] Kumar, S., & Thakur, A. (2022). Chatbots in Libraries: Enhancing User Experience."
    )
    parts = ReferenceSplitter.split_manuscript_references(raw_blob)
    assert len(parts) == 2
    assert parts[0].startswith("[1]")
    assert parts[1].startswith("[22]")
    # Verify illegal characters were stripped
    assert "\x0c" not in parts[0]
    assert "\x00" not in parts[1]


def test_excel_writer_illegal_character_resilience(tmp_path):
    from src.pipeline.custom_runner import _write_custom_excel
    import openpyxl

    bad_text = "Analysis with control char \x0c form feed and null \x00 in text [1]."
    results = [{
        "reference_no": 1,
        "reference_id": "CUSTOM_001",
        "scopus_link_valid": True,
        "scopus_reference_link": "http://example.com",
        "final_status": "VALID_CORRECT",
        "needs_human_review": False,
        "normalized_doi": "10.1000/182",
        "doi_resolves": True,
        "http_status": 200,
        "resolved_title": f"Title with \x08 bad char",
        "resolved_authors": f"Author with \x1b bad char",
        "resolved_journal": "Journal of AI",
        "resolved_year": 2023,
        "composite_score": 0.95,
        "raw_reference": bad_text,
        "decision_rationale": f"Rationale with \x0c form feed"
    }]
    stats = {
        "total_references": 1,
        "scopus_linked": 1,
        "scopus_unlinked": 0,
        "resolving_dois": 1,
        "recovered_dois": 0,
        "needs_review": 0,
        "scopus_linked_pct": 100.0,
        "resolving_pct": 100.0
    }
    source_info = {
        "source_eid": "DOCX",
        "source_title": "Paper with \x0c control char",
        "file_type": "docx",
        "page_count": 1,
        "word_count": 50
    }
    diagnostics = {
        "citation_style_display": "IEEE [1]",
        "consistency_score": 100.0,
        "orphan_count": 0,
        "missing_count": 0,
        "in_text_citations": [{
            "marker": "[1]",
            "type": "NUMERIC_BRACKETED",
            "context": "Context with \x0c control char"
        }]
    }

    excel_file = tmp_path / "test_resilience.xlsx"
    # This must complete with zero IllegalCharacterError
    _write_custom_excel(results, stats, source_info, excel_file, document_diagnostics=diagnostics)
    assert excel_file.exists()

    wb = openpyxl.load_workbook(excel_file)
    assert "Summary" in wb.sheetnames
    assert "Citation Style & In-Text" in wb.sheetnames


def test_bracket_spacing_normalization_and_splitting():
    """
    Tests that PDF font kerning artifacts like [1 0], [15 ], and [ 18 ]
    are properly normalized to canonical brackets and split into distinct references.
    """
    raw_blob = (
        "[1] First Reference (2020). First paper.\n"
        "[2] Second Reference (2021). Second paper.\n"
        "[1 0] Tenth Reference (2022). Tenth paper with spaced digits.\n"
        "[15 ] Fifteenth Reference (2023). Fifteenth paper with trailing space.\n"
        "[ 18 ] Eighteenth Reference (2024). Eighteenth paper with outer spaces.\n"
    )
    parts = ReferenceSplitter.split_manuscript_references(raw_blob)
    assert len(parts) == 5
    assert parts[0].startswith("[1]")
    assert parts[1].startswith("[2]")
    assert parts[2].startswith("[10]")
    assert parts[3].startswith("[15]")
    assert parts[4].startswith("[18]")


def test_ocr_tolerant_doi_extractor():
    """
    Tests recovery of DOIs with OCR letter substitutions and typography kerning spaces.
    """
    from src.parsing.doi_extractor import extract_doi

    cases = [
        ("doi: 10. 1 080/1 533 290X.2021 .18 813 25.", "10.1080/1533290X.2021.1881325"),
        ("doi:l0 .5430/ijlis.v7n2pl2.", "10.5430/ijlis.v7n2p12"),
        ("doi:l0.1 108/LM-04-2021-0055.", "10.1108/LM-04-2021-0055"),
        ("doi:l0 .1016 /j.acalib.2020 .10 21 10.", "10.1016/j.acalib.2020.102110"),
        ("doi: 10.55 39/jlis.v4nl p41 .", "10.5539/jlis.v4n1p41"),
        ("doi: 10.21597 /ijism.2023.3 144. 191 Authorized licensed...", "10.21597/ijism.2023.3144"),
    ]
    for raw, expected in cases:
        extracted = extract_doi(raw)
        assert extracted is not None, f"Failed to extract DOI from {raw}"
        assert extracted.lower() == expected.lower(), f"Expected {expected}, got {extracted}"


def test_citation_parser_elsevier_apa_quoted():
    """
    Tests decomposition of Elsevier numbered, APA author-year, and quoted reference strings.
    """
    from src.parsing.citation_parser import CitationParser

    # 1. Elsevier numbered style
    c1 = "[2] Z. Guo, A. Lai, J.H. Thygesen, J. Farrington, T. Keen, K. Li, Large language models for mental health applications: systematic review, JMIR Ment. Health 11 (2024) e57400."
    p1 = CitationParser.parse(c1)
    assert "Large language models" in p1["cited_title"]
    assert "Guo" in p1["cited_authors"]
    assert p1["cited_year"] == 2024
    assert p1["cited_volume"] == "11"
    assert p1["cited_pages"] == "e57400"

    # 2. APA author-year style
    c2 = "Vaswani, A., Shazeer, N., & Parmar, N. (2017). Attention is all you need. Advances in Neural Information Processing Systems, 30, 5998-6008."
    p2 = CitationParser.parse(c2)
    assert p2["cited_title"] == "Attention is all you need"
    assert "Vaswani" in p2["cited_authors"]
    assert p2["cited_year"] == 2017

    # 3. Quoted title style
    c3 = '[1] Kumar, R., & Singh, B. "Usability of OPAC Systems", Journal of Information Science, 47(4), 520-533.'
    p3 = CitationParser.parse(c3)
    assert p3["cited_title"] == "Usability of OPAC Systems"
    assert "Kumar" in p3["cited_authors"]


def test_extract_document_doi_with_unicode_ligatures():
    """
    Tests manuscript-level DOI detection with unicode ligatures (ff, fi) and link annotations.
    """
    from src.parsing.doi_extractor import extract_doi, extract_document_doi

    # Text containing unicode ligature \ufb00 (LATIN SMALL LIGATURE FF)
    ligature_text = "https://doi.org/10.1016/j.in\ufb00us.2026.104702\njournal homepage: www.elsevier.com"
    doi_extracted = extract_doi(ligature_text)
    assert doi_extracted == "10.1016/j.inffus.2026.104702"

    links = ["https://doi.org/10.1016/j.inffus.2026.104702"]
    doc_doi = extract_document_doi("Random text", links=links)
    assert doc_doi == "10.1016/j.inffus.2026.104702"


def test_author_date_manuscript_splitting():
    """
    Tests unnumbered Author-Date reference splitting (e.g. Taylor & Francis, APA, Harvard).
    """
    blob = (
        "Barba, Ian, Ryan Cassidy, and Esther De Leon. 2013. “Web analytics reveal user behavior: TTU libraries’ experience with Google Analytics.” Journal of Web Librarianship 7 (4):389–400. http://dx.doi.org/10.1080/19322909.2013.82899.\n"
        "Betty, Paul. 2009. “Assessing homegrown library collections: Using Google Analytics.” Journal of Electronic Resources Librarianship 21 (1):75. http://dx.doi.org/10.1080/19411260902858631.\n"
        "Google. n.d.a. About the flow visualization reports. https://support.google.com/analytics/answer/2519986.\n"
        "Loftus, Wayne. 2012. “Demonstrating success: Web analytics.” Journal of Web Librarianship 6 (1):45–55. http://dx.doi.org/10.1080/19322909.2012.651416.\n"
    )
    parts = ReferenceSplitter.split_manuscript_references(blob)
    assert len(parts) == 4
    assert "Barba" in parts[0]
    assert "Betty" in parts[1]
    assert "Google" in parts[2]
    assert "Loftus" in parts[3]


def test_author_date_citation_style_detection():
    """
    Tests Author-Year in-text citation scanning with and without commas, and page numbers.
    """
    body = (
        "Several studies evaluated web analytics in libraries (Barba et al. 2013, 392). "
        "Turner (2010) described key performance indicators. "
        "Other authors also noted that user satisfaction increases (Fagan 2014; Loftus 2012, 47). "
        "Google provides event tracking documentation (Google n.d.a)."
    )
    style, in_text = CitationStyleDetector.scan_in_text_citations(body)
    assert style == CitationStyleFamily.AUTHOR_YEAR
    assert len(in_text) >= 4


if __name__ == "__main__":
    pytest.main(["-v", __file__])

